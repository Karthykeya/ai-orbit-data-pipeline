import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

mcp = json.load(open("data/mcp_servers.json"))
companies = json.load(open("data/companies.json"))
repositories_path = "data/repositories.json"
try:
    repositories = json.load(open(repositories_path))
except FileNotFoundError:
    repositories = []
rels = json.load(open("data/relationships.json"))

wb = Workbook()

# ---------- Sheet 1: MCP Servers (main dataset) ----------
ws = wb.active
ws.title = "MCP Servers"
headers = ["id", "entity_type", "name", "description", "description_source",
           "official_url", "logo_url", "categories", "vendor", "server_url",
           "auth_type", "transport", "installation", "verification_status",
           "source_name", "source_url"]
ws.append(headers)
for rec in mcp:
    ws.append([
        rec["id"], rec["entity_type"], rec["name"], rec["description"],
        rec.get("description_source", ""),
        rec["url"], rec["logo_url"], ", ".join(rec["categories"]),
        rec["metadata"]["vendor"], rec["metadata"]["server_url"],
        rec["metadata"]["auth_type"], rec["metadata"]["transport"],
        rec["metadata"]["installation"], rec["verification_status"],
        rec["source"]["name"], rec["source"]["url"],
    ])

# ---------- Sheet 2: Companies ----------
ws2 = wb.create_sheet("Companies")
headers2 = ["id", "entity_type", "name", "description", "official_url",
            "categories", "headquarters", "founding_year", "industry_sector",
            "source_name", "source_url"]
ws2.append(headers2)
for rec in companies:
    ws2.append([
        rec["id"], rec["entity_type"], rec["name"], rec["description"], rec["url"],
        ", ".join(rec["categories"]), rec.get("headquarters") or "",
        rec.get("founding_year") or "", rec["industry_sector"],
        rec["source"]["name"], rec["source"]["url"],
    ])

# ---------- Sheet 2b: Repositories (community tier) ----------
ws2b = wb.create_sheet("Repositories")
headers2b = ["id", "entity_type", "name", "description", "official_url",
             "categories", "owner", "stars", "primary_language", "last_updated",
             "source_name", "source_url"]
ws2b.append(headers2b)
for rec in repositories:
    ws2b.append([
        rec["id"], rec["entity_type"], rec["name"], rec["description"], rec["url"],
        ", ".join(rec["categories"]), rec.get("owner") or "",
        rec.get("stars") if rec.get("stars") is not None else "",
        rec.get("primary_language") or "", rec.get("last_updated") or "",
        rec["source"]["name"], rec["source"]["url"],
    ])

# ---------- Sheet 3: Relationships ----------
ws3 = wb.create_sheet("Relationships")
headers3 = ["id", "type", "source_id", "source_type", "target_id", "target_type", "note"]
ws3.append(headers3)
for rec in rels:
    ws3.append([
        rec["id"], rec["type"], rec["source_id"], rec["source_type"],
        rec["target_id"], rec["target_type"], rec.get("note", ""),
    ])

# ---------- Sheet 4: Quality Report ----------
quality = json.load(open("data/quality_report.json"))
ws_q = wb.create_sheet("Quality Report")
ws_q.append(["Metric", "Value"])
flat_rows = [
    ("MCP records (total)", quality["totals"]["mcp_records"]),
    ("  - verified (first-party)", quality["totals"]["verified_first_party_records"]),
    ("  - community_github_verified", quality["totals"]["community_github_verified_records"]),
    ("Company records", quality["totals"]["company_records"]),
    ("Repository records", quality["totals"]["repository_records"]),
    ("Relationship records", quality["totals"]["relationship_records"]),
    ("Duplicate ids", quality["integrity"]["duplicate_ids"]),
    ("Duplicate name+vendor pairs", quality["integrity"]["duplicate_name_vendor_pairs"]),
    ("Passes uniqueness check", quality["integrity"]["passes_uniqueness_check"]),
    ("Description coverage %", quality["completeness"]["description_coverage_pct"]),
    ("Official URL coverage %", quality["completeness"]["official_url_coverage_pct"]),
    ("Logo URL coverage %", quality["completeness"]["logo_url_coverage_pct"]),
    ("Verified tier: URL & source same host %", quality["traceability"]["verified_tier_url_and_source_same_host_pct"]),
    ("Community tier: URL matches declared repo %", quality["traceability"]["community_tier_url_matches_declared_repo_pct"]),
    ("Companies with HQ %", quality["company_enrichment"]["companies_with_headquarters_pct"]),
    ("Companies with founding year %", quality["company_enrichment"]["companies_with_founding_year_pct"]),
    ("Repositories with GitHub star count %", quality["repository_enrichment"]["repositories_with_github_stars_pct"]),
]
for label, value in flat_rows:
    ws_q.append([label, value if value is not None else "n/a"])
ws_q.append([])
ws_q.append(["Description sourcing breakdown:"])
for source, count in quality["description_sourcing"].items():
    ws_q.append([f"  {source}", count])
ws_q.append([])
ws_q.append(["Regenerate this report anytime with: python scripts/quality_report.py"])

# ---------- Sheet 5: README ----------
ws4 = wb.create_sheet("README", 0)
t = quality["totals"]
verified_n = t["verified_first_party_records"]
community_n = t["community_github_verified_records"]
total_n = t["mcp_records"]
readme_lines = [
    ("AI Orbit Data Ingestion — MCP Servers Module", True),
    ("", False),
    (f"Scope: {total_n} MCP (Model Context Protocol) servers —", False),
    (f"  {verified_n} first-party (verified against each vendor's own documentation)", False),
    (f"  {community_n} community (from Anthropic's official MCP Registry, each with a real public GitHub repo)", False),
    ("See verification_status per row in the MCP Servers tab for which is which.", False),
    ("", False),
    ("Sheets:", True),
    (f"  MCP Servers     - main dataset ({total_n} records: {verified_n} verified + {community_n} community)", False),
    (f"  Companies       - derived vendor/company entities ({len(companies)} records, verified tier only),", False),
    ("                    enriched with real headquarters, founding year, and industry sector", False),
    (f"  Repositories    - derived GitHub repository entities ({len(repositories)} records, community tier only)", False),
    (f"  Relationships   - develops/integrates_with (verified) + hosted_in (community) edges ({t['relationship_records']})", False),
    ("  Quality Report  - automated, regenerable integrity/completeness metrics", False),
    ("", False),
    ("Quality snapshot (see Quality Report tab for the live numbers):", True),
    (f"  - {quality['integrity']['duplicate_ids']} duplicate ids, {quality['integrity']['duplicate_name_vendor_pairs']} duplicate name+vendor pairs", False),
    (f"  - {quality['completeness']['description_coverage_pct']}% description / {quality['completeness']['official_url_coverage_pct']}% official-URL / {quality['completeness']['logo_url_coverage_pct']}% logo-URL coverage", False),
    (f"  - Description sourcing: {quality['description_sourcing']}", False),
    ("", False),
    ("Data quality notes:", True),
    ("  - verified tier: url = official documentation page for each server (not a 3rd-party directory)", False),
    ("  - community tier: url = the server's own public GitHub repository, sourced via the official MCP Registry", False),
    ("  - logo_url = Clearbit Logo API (verified tier) or the GitHub owner's own avatar (community tier)", False),
    ("  - descriptions: LLM-authored for the verified tier; publisher-provided (from the registry) for community", False),
    ("    records unless run with an ANTHROPIC_API_KEY set, in which case they're also LLM-rewritten live", False),
    ("  - ids are deterministic UUIDv5s so re-running the pipeline is idempotent", False),
    ("", False),
    ("Pipeline, tests, CI, and ARCHITECTURE.md: see the accompanying GitHub repository.", False),
    ("This tab is generated automatically by build_sheet.py from the actual data — it will never drift out of sync.", False),
]
for i, (text, bold) in enumerate(readme_lines, start=1):
    cell = ws4.cell(row=i, column=1, value=text)
    cell.font = Font(name="Arial", bold=bold, size=13 if bold and i == 1 else 11)

# ---------- Formatting ----------
header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
header_font = Font(name="Arial", bold=True, color="FFFFFF")
for sheet in (ws, ws2, ws2b, ws3, ws_q):
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    for col_cells in sheet.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        sheet.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=False)

ws4.column_dimensions["A"].width = 90

wb.save("data/AI_Orbit_MCP_Dataset.xlsx")
print("saved")
